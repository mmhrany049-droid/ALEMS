"""Phase 6 — Exam Center & Analytics & Report & Export (doc 12، doc 08 §8.1، doc 06).

پذیرش کلیدی: V2-A01 (overview سه متریک جدا — هرگز یک عدد قاطی) و
V2-A02 (export json شامل کتاب و attempt).
"""
from __future__ import annotations

import io

from tests.test_planning import (
    J_SATURDAY, J_TODAY, J_TOMORROW, PASS, SATURDAY_ISO, TODAY_ISO, TOMORROW_ISO,
    _book, _h, _user,
)  # noqa: F401 — helpers مشترک



def _setup(client, email, taught=True):
    """user + book(۱۰ سوال) + taught — بدون جلسه تست."""
    u = _user(client, email)
    r = client.post("/api/v1/resources/import-book", json=_book(), headers=_h(u))
    rid = r.json()["data"]["resource"]["id"]
    tree = client.get(f"/api/v1/resources/{rid}/tree", headers=_h(u)).json()["data"]
    topic_id = tree["tree"][0]["children"][0]["id"]
    if taught:
        client.put(
            "/api/v1/students/me/taught-topics",
            json={"items": [{"topic_id": topic_id, "taught": True}]},
            headers=_h(u),
        )
    return u, rid, topic_id


def _session(client, u, rid, correct=(), wrong=(), finish=True):
    """یک جلسه تست با پاسخ‌های مشخص → dict جلسه (تمام‌شده)."""
    s = client.post("/api/v1/test-sessions", json={"resource_id": rid}, headers=_h(u)).json()["data"]
    qid = {q["number"]: q["question_id"] for q in s["questions"]}
    items = [{"question_id": qid[n], "status": "answered", "result": "correct"} for n in correct]
    items += [{"question_id": qid[n], "status": "answered", "result": "wrong"} for n in wrong]
    if items:
        r = client.post(f"/api/v1/test-sessions/{s['session']['id']}/records", json={"items": items}, headers=_h(u))
        assert r.status_code == 200, r.text
    if finish:
        r = client.post(f"/api/v1/test-sessions/{s['session']['id']}/finish", json={}, headers=_h(u))
        assert r.status_code == 200, r.text
        return r.json()["data"]["session"]
    return s["session"]


# --- Exam Center (doc 12 §12.2) -----------------------------------------------------------

def test_exam_lifecycle_with_session(client):
    u, rid, topic_id = _setup(client, "ex1@example.com")
    s = _session(client, u, rid, correct=(1, 2), wrong=(3, 4, 5))

    r = client.post("/api/v1/exams", headers=_h(u), json={
        "title": "آزمون آزمایشی ۱", "kind": "mock",
        "scheduled_date": J_TOMORROW, "planned_duration_minutes": 90,
        "subjects": ["ریاضی"], "planned_topic_ids": [topic_id],
    })
    assert r.status_code == 200, r.text
    e = r.json()["data"]
    assert e["status"] == "planned" and e["status_fa"] == "برنامه‌ریزی‌شده"
    assert e["kind_fa"] == "آزمایشی"
    assert e["scheduled_date"] == TOMORROW_ISO          # شمسی در ورودی قبول است
    assert e["scheduled_date_jalali"] == J_TOMORROW

    lst = client.get("/api/v1/exams", headers=_h(u)).json()["data"]
    assert len(lst["items"]) == 1 and lst["upcoming_count"] == 1

    st = client.post(f"/api/v1/exams/{e['id']}/start", headers=_h(u))
    assert st.json()["data"]["status"] == "in_progress"
    st2 = client.post(f"/api/v1/exams/{e['id']}/start", headers=_h(u))
    assert st2.status_code == 422  # شروع دوباره → انتقال نامعتبر

    sub = client.post(f"/api/v1/exams/{e['id']}/submit", headers=_h(u), json={"session_ids": [s["id"]]})
    assert sub.status_code == 200, sub.text
    d = sub.json()["data"]
    assert d["status"] == "finished" and d["finished_at"]
    sc = d["scoring"]
    # T=10 (کل سوالات جلسه) C=2 W=3 U=5 — §8.1 با k=0.33 از settings
    assert sc["total_count"] == 10 and sc["correct_count"] == 2 and sc["wrong_count"] == 3
    assert sc["percent_konkur"] == round((2 - 0.33 * 3) / 10 * 100, 2)   # 10.1
    assert sc["percent_no_penalty"] == 20.0
    assert sc["percent_konkur"] != sc["percent_no_penalty"]              # دو درصد همیشه جدا
    assert topic_id in d["actual_topic_ids"]                              # doc 12 §12.2

    res = client.get(f"/api/v1/exams/{e['id']}/result", headers=_h(u)).json()["data"]
    assert len(res["result"]["sessions"]) == 1
    assert res["result"]["sessions"][0]["percent_no_penalty"] == 20.0
    assert len(res["result"]["planned_topics"]) == 1

    # سند تمام‌شده ویرایش نمی‌شود؛ submit دوباره هم نامعتبر است
    ed = client.put(f"/api/v1/exams/{e['id']}", headers=_h(u), json={"title": "عوض شد"})
    assert ed.status_code == 422 and "قابل ویرایش نیست" in ed.json()["error"]["message"]
    sub2 = client.post(f"/api/v1/exams/{e['id']}/submit", headers=_h(u), json={"session_ids": [s["id"]]})
    assert sub2.status_code == 422


def test_exam_manual_counts_negative_and_cancel(client):
    u, rid, _t = _setup(client, "ex2@example.com")

    # شمارش دستی — امتحان درسی بدون جلسه دیجیتال
    r = client.post("/api/v1/exams", headers=_h(u), json={
        "title": "امتحان ریاضی مدرسه", "kind": "school_subject", "resource_id": rid,
    })
    e = r.json()["data"]
    sub = client.post(f"/api/v1/exams/{e['id']}/submit", headers=_h(u), json={
        "total_count": 20, "correct_count": 10, "wrong_count": 4, "actual_duration_minutes": 70,
    })
    sc = sub.json()["data"]["scoring"]
    assert sc["unanswered_count"] == 6
    assert sc["percent_konkur"] == round((10 - 0.33 * 4) / 20 * 100, 2)  # 43.4
    assert sc["percent_no_penalty"] == 50.0
    assert sc["percent_konkur"] < sc["percent_no_penalty"]
    assert sub.json()["data"]["actual_duration_seconds"] == 70 * 60

    # درصد منفی نمایش داده می‌شود (§8.1 show_negative)
    r2 = client.post("/api/v1/exams", headers=_h(u), json={"title": "آزمون سخت", "kind": "free"})
    e2 = r2.json()["data"]
    sub2 = client.post(f"/api/v1/exams/{e2['id']}/submit", headers=_h(u), json={
        "total_count": 10, "correct_count": 0, "wrong_count": 9,
    })
    assert sub2.json()["data"]["scoring"]["percent_konkur"] < 0

    # cancel و بازگشت
    r3 = client.post("/api/v1/exams", headers=_h(u), json={"title": "لغو شدنی"})
    e3 = r3.json()["data"]
    can = client.put(f"/api/v1/exams/{e3['id']}", headers=_h(u), json={"status": "cancelled"})
    assert can.json()["data"]["status"] == "cancelled" and can.json()["data"]["status_fa"] == "لغوشده"
    bad = client.post(f"/api/v1/exams/{e3['id']}/start", headers=_h(u))
    assert bad.status_code == 422

    # submit خالی
    r4 = client.post("/api/v1/exams", headers=_h(u), json={"title": "خالی"})
    empty = client.post(f"/api/v1/exams/{r4.json()['data']['id']}/submit", headers=_h(u), json={})
    assert empty.status_code == 422 and "جلسه آزمون را انتخاب کن" in empty.json()["error"]["message"]


def test_exam_multi_session_merge(client):
    """mock چنددرس — جمع شمارش چند جلسه (doc 12 §12.2 subjects[])."""
    u, rid, _t = _setup(client, "ex3@example.com")
    s1 = _session(client, u, rid, correct=(1, 2), wrong=(3, 4, 5))
    s2 = _session(client, u, rid, correct=(6, 7, 8, 9), wrong=(10,))
    r = client.post("/api/v1/exams", headers=_h(u), json={
        "title": "آزمایشی جامع", "kind": "mock", "subjects": ["ریاضی", "فیزیک"],
    })
    e = r.json()["data"]
    sub = client.post(f"/api/v1/exams/{e['id']}/submit", headers=_h(u), json={"session_ids": [s1["id"], s2["id"]]})
    sc = sub.json()["data"]["scoring"]
    assert sc["total_count"] == 20 and sc["correct_count"] == 6 and sc["wrong_count"] == 4
    assert sc["percent_konkur"] == round((6 - 0.33 * 4) / 20 * 100, 2)  # 23.4
    assert sc["percent_no_penalty"] == 30.0
    assert len(sub.json()["data"]["session_ids"]) == 2


def test_today_upcoming_exam_and_priority_boost(client):
    """آزمون نزدیک در Today Hub (doc 07.6 بخش ۵) + boost اولویت (doc 11.3 مرحله ۲)."""
    u, rid, topic_id = _setup(client, "ex4@example.com")
    _session(client, u, rid, wrong=(1, 2, 3))  # learning state + review demand

    client.post("/api/v1/exams", headers=_h(u), json={
        "title": "آزمون فردا", "kind": "mock", "scheduled_date": J_TOMORROW,
        "planned_topic_ids": [topic_id],
    })

    t = client.get("/api/v1/today", headers=_h(u)).json()["data"]
    assert len(t["upcoming_exams"]) == 1
    ex = t["upcoming_exams"][0]
    assert ex["title"] == "آزمون فردا" and ex["days_until"] == 1 and ex["kind_fa"] == "آزمایشی"

    g = client.post("/api/v1/plans/generate-week", headers=_h(u), json={}).json()["data"]
    assert g["steps"][1]["step"] == "exams" and g["steps"][1]["info"]["upcoming"] == 1
    pr = [p for p in g["priority"] if p["topic_id"] == topic_id][0]
    assert "exam_prep" in pr["reason_codes"]


# --- Analytics (doc 12 §12.1، V2-A01) -------------------------------------------------------

def test_analytics_overview_three_separate_blocks_v2_a01(client):
    u, rid, topic_id = _setup(client, "ex5@example.com")
    _session(client, u, rid, correct=(1, 2), wrong=(3, 4, 5))

    d = client.get("/api/v1/analytics/overview", headers=_h(u)).json()["data"]

    # V2-A01 — سه بلوک مستقل، هرگز یک عدد قاطی
    assert isinstance(d["coverage"], dict) and isinstance(d["accuracy"], dict) and isinstance(d["volume"], dict)
    assert d["coverage"]["topics_attempted"] == 1 and d["coverage"]["topics_total"] == 1
    assert d["coverage"]["questions_attempted"] == 5 and d["coverage"]["questions_total"] == 10
    assert d["coverage"]["topics_ratio"] == 1.0 and d["coverage"]["questions_ratio"] == 0.5

    assert d["accuracy"]["correct"] == 2 and d["accuracy"]["wrong"] == 3
    assert d["accuracy"]["answered_accuracy"] == 0.4
    assert d["accuracy"]["percent_konkur"] == round((2 - 0.33 * 3) / 10 * 100, 2)  # 10.1 — T=کل سوالات جلسه
    assert d["accuracy"]["percent_no_penalty"] == 20.0
    assert d["accuracy"]["percent_konkur"] != d["accuracy"]["percent_no_penalty"]

    assert d["volume"]["attempts"] == 5 and d["volume"]["sessions"] == 1
    assert d["volume"]["active_days"] >= 1

    assert len(d["daily"]) == 30
    assert len(d["time_buckets"]) >= 1 and all("answered_accuracy" in b for b in d["time_buckets"])
    assert d["window"]["days"] == 30
    # بدون هدف کنکور → پیام ثبت هدف (doc 12 §12.4 — بدون ادعای رتبه)
    assert d["konkurs_target"]["has_target"] is False
    assert "رتبه" not in d["konkurs_target"]["message_fa"]


def test_analytics_target_tracker_qualitative(client):
    u, rid, _t = _setup(client, "ex6@example.com")
    client.put("/api/v1/students/me", headers=_h(u), json={"target": "رشته تجربی"})
    _session(client, u, rid, correct=(1, 2), wrong=(3, 4, 5))
    d = client.get("/api/v1/analytics/overview", headers=_h(u)).json()["data"]
    tgt = d["konkurs_target"]
    assert tgt["has_target"] is True and tgt["target"] == "رشته تجربی"
    assert tgt["level_fa"] and tgt["message_fa"]
    assert tgt["based_on"]["answered_accuracy"] == 0.4
    # فقط کیفی — هیچ عدد/رتبه‌ای ادعا نمی‌شود
    assert not any(ch.isdigit() for ch in tgt["level_fa"])


def test_analytics_slices(client):
    u, rid, topic_id = _setup(client, "ex7@example.com")
    _session(client, u, rid, correct=(1, 2), wrong=(3, 4, 5))

    subj = client.get("/api/v1/analytics/by-subject", headers=_h(u)).json()["data"]["items"]
    assert len(subj) == 1
    row = subj[0]
    assert row["title"] == "ریاضی جامع"
    assert row["coverage"]["topics_ratio"] == 1.0
    assert row["accuracy"]["answered_accuracy"] == 0.4
    assert row["volume"]["attempts"] == 5

    chap = client.get("/api/v1/analytics/by-chapter", headers=_h(u)).json()["data"]["items"]
    assert chap[0]["chapter_title"] == "فصل ۱" and chap[0]["coverage"]["topics_total"] == 1

    top = client.get("/api/v1/analytics/by-topic", headers=_h(u)).json()["data"]["items"]
    assert top[0]["topic_title"] == "تابع"
    assert top[0]["coverage"]["questions_attempted"] == 5
    assert top[0]["exam_readiness"] is not None  # learning state بعد از finish ساخته شده

    top_acc = client.get("/api/v1/analytics/by-topic?order=accuracy", headers=_h(u)).json()["data"]["items"]
    assert top_acc[0]["accuracy"]["answered_accuracy"] == 0.4

    diff = client.get("/api/v1/analytics/difficulty", headers=_h(u)).json()["data"]["items"]
    assert diff[0]["difficulty_fa"] == "نامشخص" and diff[0]["volume"]["attempts"] == 5

    mis = client.get("/api/v1/analytics/mistakes", headers=_h(u)).json()["data"]
    assert mis["top_wrong_topics"][0]["topic_title"] == "تابع"
    assert mis["top_wrong_topics"][0]["wrong_count"] == 3
    assert mis["notes_total"] == 3 and mis["notes_unlabeled"] == 3
    assert mis["by_error_type"][0]["error_type_fa"] == "بدون برچسب"


# --- Reports (doc 12 §12.5) -------------------------------------------------------------------

def test_reports_daily_weekly_monthly(client):
    u, rid, _t = _setup(client, "ex8@example.com")
    _session(client, u, rid, correct=(1, 2), wrong=(3, 4, 5))
    day = client.put(f"/api/v1/plans/{TODAY_ISO}", headers=_h(u),
                     json={"tasks": [{"title": "مطالعه تابع", "minutes": 45}]}).json()["data"]
    client.post(f"/api/v1/plans/tasks/{day['tasks'][0]['id']}/status", headers=_h(u), json={"status": "done"})

    d = client.get("/api/v1/reports/daily", headers=_h(u)).json()["data"]
    assert d["kind"] == "daily" and d["date"] == TODAY_ISO
    assert d["volume"]["study_minutes"] == 45
    assert d["accuracy"]["correct"] == 2 and d["accuracy"]["wrong"] == 3
    assert d["plan"]["tasks_done"] == 1 and d["plan"]["tasks_total"] == 1
    assert "coverage" in d and "accuracy" in d and "volume" in d  # سه بلوک جدا در گزارش

    dj = client.get(f"/api/v1/reports/daily?date={J_TODAY}", headers=_h(u)).json()["data"]
    assert dj["date"] == TODAY_ISO  # شمسی قبول است

    w = client.get("/api/v1/reports/weekly", headers=_h(u)).json()["data"]
    assert w["kind"] == "weekly" and w["week_start"] == SATURDAY_ISO and len(w["days"]) == 7
    assert w["volume"]["attempts"] == 5
    assert w["previous_week"]["attempts_delta"] == 5
    assert "used_minutes" in w["capacity"]

    m = client.get("/api/v1/reports/monthly?month=1405/06", headers=_h(u)).json()["data"]
    assert m["kind"] == "monthly" and m["month"] == "1405/06"
    assert m["start"] == "2026-08-23" and m["start_jalali"] == "1405/06/01"
    assert len(m["weeks"]) >= 4

    m2 = client.get("/api/v1/reports/monthly?month=2026-09", headers=_h(u)).json()["data"]
    assert m2["start"] == "2026-09-01" and m2["end"] == "2026-09-30"

    bad = client.get("/api/v1/reports/monthly?month=abc", headers=_h(u))
    assert bad.status_code == 422 and "ماه را مثل" in bad.json()["error"]["message"]


# --- Export (doc 12 §12.5، V2-A02) ---------------------------------------------------------------

def test_export_json_full_v2_a02(client):
    u, rid, topic_id = _setup(client, "ex9@example.com")
    _session(client, u, rid, correct=(1,), wrong=(2, 3))
    client.post("/api/v1/exams", headers=_h(u), json={"title": "آزمون من", "kind": "mock"})

    d = client.get("/api/v1/export/json", headers=_h(u)).json()["data"]
    assert d["app"] == "ALEMS" and d["exported_at_jalali"] == J_TODAY
    assert d["profile"]["email"] == "ex9@example.com"
    assert d["settings"]["konkurs_penalty_k"] == 0.33
    # V2-A02 — کتاب‌ها و attemptها
    assert len(d["books"]) == 1
    b = d["books"][0]
    assert b["title"] == "ریاضی جامع"
    ch = b["tree"][0]
    assert ch["title"] == "فصل ۱" and ch["children"][0]["title"] == "تابع"
    assert len(ch["children"][0]["questions"]) == 10
    assert len(d["attempts"]) == 3
    assert {a["result"] for a in d["attempts"]} == {"correct", "wrong"}
    assert len(d["test_sessions"]) == 1
    assert len(d["exams"]) == 1 and d["exams"][0]["title"] == "آزمون من"
    assert d["profile"]["taught_topics"][0]["topic_id"] == topic_id


def test_export_excel_rtl(client):
    from openpyxl import load_workbook

    u, rid, _t = _setup(client, "ex10@example.com")
    _session(client, u, rid, correct=(1,), wrong=(2, 3, 4))

    r = client.get("/api/v1/export/excel", headers=_h(u))
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"  # zip/xlsx magic

    wb = load_workbook(io.BytesIO(r.content))
    names = wb.sheetnames
    for want in ("خلاصه", "کتاب‌ها", "تست‌ها", "تلاش‌ها", "مرور", "برنامه", "آزمون‌ها"):
        assert want in names, names
    ws = wb["خلاصه"]
    assert ws.sheet_view.rightToLeft is True
    texts = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "پوشش" in texts and "دقت" in texts and "حجم" in texts  # سه بلوک جدا در خلاصه
    att = wb["تلاش‌ها"]
    assert att.max_row == 5  # header + 4 attempt


def test_export_pdf_rtl(client):
    u, rid, _t = _setup(client, "ex11@example.com")
    _session(client, u, rid, correct=(1,), wrong=(2,))
    client.post("/api/v1/exams", headers=_h(u), json={
        "title": "آزمون هفته", "kind": "mock", "scheduled_date": TOMORROW_ISO,
    })

    for kind in ("weekly", "daily", "monthly", "summary"):
        r = client.get(f"/api/v1/export/pdf?report={kind}", headers=_h(u))
        assert r.status_code == 200, (kind, r.text[:200])
        assert r.headers["content-type"] == "application/pdf"
        assert r.content[:5] == b"%PDF-"
        assert b"Vazirmatn" in r.content  # فونت فارسی embed شده
        assert len(r.content) > 2500

    # weekly با week_start شمسی
    r = client.get(f"/api/v1/export/pdf?report=weekly&when={J_SATURDAY}", headers=_h(u))
    assert r.status_code == 200 and r.content[:5] == b"%PDF-"

    bad = client.get("/api/v1/export/pdf?report=nope", headers=_h(u))
    assert bad.status_code == 422


# --- خطاهای فارسی -----------------------------------------------------------------------------

def test_exam_persian_errors(client):
    u, _rid, _t = _setup(client, "ex12@example.com")

    nf = client.get("/api/v1/exams/nope", headers=_h(u))
    assert nf.status_code == 404 and nf.json()["error"]["message"] == "آزمون پیدا نشد."

    empty_title = client.post("/api/v1/exams", headers=_h(u), json={"title": "   "})
    assert empty_title.status_code == 422
    assert empty_title.json()["error"]["message"] == "عنوان آزمون نمی‌تواند خالی باشد."

    bad_res = client.post("/api/v1/exams", headers=_h(u), json={"title": "x", "resource_id": "nope"})
    assert bad_res.status_code == 404 and bad_res.json()["error"]["message"] == "کتاب پیدا نشد."

    r = client.post("/api/v1/exams", headers=_h(u), json={"title": "موقت"})
    bad_sess = client.post(f"/api/v1/exams/{r.json()['data']['id']}/submit", headers=_h(u),
                           json={"session_ids": ["nope"]})
    assert bad_sess.status_code == 404 and bad_sess.json()["error"]["message"] == "جلسه آزمون پیدا نشد."

    unauth = client.get("/api/v1/analytics/overview")
    assert unauth.status_code == 401 and unauth.json()["error"]["message"] == "برای ادامه باید وارد شوید."
