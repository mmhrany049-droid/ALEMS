"""Export — JSON کامل کاربر / Excel RTL / PDF فارسی (doc 12 §12.5، V2-A02).

- JSON: پروفایل، کتاب‌ها (فهرست کامل)، sessionها، **همه attemptها**، مرور،
  learning states، برنامه، اهداف، آزمون‌ها، تنظیمات.
- Excel: openpyxl با sheet_view.rightToLeft و سرستون فارسی.
- PDF: reportlab + Vazirmatn TTF (embedding) + arabic_reshaper + python-bidi → RTL واقعی.
"""
from __future__ import annotations

import datetime as dt
import io
import logging
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.jalali import gregorian_to_jalali
from app.modules.academic.models import Question, Resource, Topic
from app.modules.activity.models import AttemptResult, ErrorNote, QuestionMark, TestSession
from app.modules.analytics.service import TZ_TEHRAN, overview as analytics_overview
from app.modules.exam.models import Exam
from app.modules.identity.models import User
from app.modules.planning.models import (
    CapacitySnapshot,
    Goal,
    PlanTask,
    Recommendation,
    TimeBlock,
)
from app.modules.report import service as report_service
from app.modules.review.models import LearningState, ReviewItem
from app.modules.settings.service import get_all as get_settings_map
from app.modules.student.models import Checkin, Student, TaughtTopic

logger = logging.getLogger("alems.export")

_FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def _num(x: Any) -> str:
    return str(x).translate(_FA_DIGITS)


def _pct(x: float | None) -> str:
    return "—" if x is None else _num(f"{round(x * 100)}٪")


def _today() -> dt.date:
    return dt.datetime.now(dt.timezone.utc).astimezone(TZ_TEHRAN).date()


def _jalali_str(d: dt.date) -> str:
    jy, jm, jd = gregorian_to_jalali(d.year, d.month, d.day)
    return f"{jy:04d}/{jm:02d}/{jd:02d}"


def _iso(x: Any) -> str | None:
    if x is None:
        return None
    if isinstance(x, dt.datetime):
        return x.isoformat()
    if isinstance(x, dt.date):
        return x.isoformat()
    return str(x)


# --- JSON کامل (V2-A02: شامل کتاب و attempt) ------------------------------------------------

def full_json(db: Session, student: Student) -> dict[str, Any]:
    user = db.execute(select(User).where(User.id == student.user_id)).scalar_one_or_none()
    resources = db.execute(select(Resource).where(Resource.student_id == student.id)).scalars().all()
    res_ids = [r.id for r in resources]
    topics = db.execute(select(Topic).where(Topic.resource_id.in_(res_ids))).scalars().all() if res_ids else []
    topic_ids = [t.id for t in topics]
    questions = db.execute(select(Question).where(Question.topic_id.in_(topic_ids))).scalars().all() if topic_ids else []
    q_by_topic: dict[str, list[Question]] = {}
    for q in questions:
        q_by_topic.setdefault(q.topic_id, []).append(q)

    def topic_out(t: Topic) -> dict[str, Any]:
        return {
            "id": t.id,
            "title": t.title,
            "block_type": t.block_type,
            "is_structural": t.is_structural,
            "questions": [
                {"id": q.id, "number": q.number, "difficulty": q.difficulty, "importance": q.importance, "tags": q.tags}
                for q in sorted(q_by_topic.get(t.id, []), key=lambda x: x.number or 0)
            ],
            "children": [topic_out(c) for c in topics if c.parent_id == t.id],
        }

    books = [
        {
            "id": r.id,
            "title": r.title,
            "publisher": r.publisher,
            "subject": r.subject,
            "created_at": _iso(r.created_at),
            "tree": [topic_out(t) for t in topics if t.resource_id == r.id and t.parent_id is None],
        }
        for r in resources
    ]

    attempts = db.execute(
        select(AttemptResult).where(AttemptResult.student_id == student.id).order_by(AttemptResult.solved_at)
    ).scalars().all()
    sessions = db.execute(select(TestSession).where(TestSession.student_id == student.id)).scalars().all()
    notes = db.execute(select(ErrorNote).where(ErrorNote.student_id == student.id)).scalars().all()
    marks = db.execute(
        select(QuestionMark).where(QuestionMark.student_id == student.id)
    ).scalars().all()
    review_items = db.execute(select(ReviewItem).where(ReviewItem.student_id == student.id)).scalars().all()
    states = db.execute(select(LearningState).where(LearningState.student_id == student.id)).scalars().all()
    goals = db.execute(select(Goal).where(Goal.student_id == student.id)).scalars().all()
    exams = db.execute(select(Exam).where(Exam.student_id == student.id)).scalars().all()
    plan_tasks = db.execute(select(PlanTask).where(PlanTask.student_id == student.id)).scalars().all()
    time_blocks = db.execute(select(TimeBlock).where(TimeBlock.student_id == student.id)).scalars().all()
    capacity = db.execute(select(CapacitySnapshot).where(CapacitySnapshot.student_id == student.id)).scalars().all()
    recs = db.execute(select(Recommendation).where(Recommendation.student_id == student.id)).scalars().all()
    checkins = db.execute(select(Checkin).where(Checkin.student_id == student.id)).scalars().all()
    taught = db.execute(select(TaughtTopic).where(TaughtTopic.student_id == student.id)).scalars().all()

    return {
        "app": "ALEMS",
        "version": get_settings().app_version,
        "exported_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "exported_at_jalali": _jalali_str(_today()),
        "profile": {
            "email": user.email if user else None,
            "full_name": user.full_name if user else None,
            "grade": student.grade,
            "track": student.track,
            "target": student.target,
            "taught_topics": [{"topic_id": t.topic_id, "taught": t.taught} for t in taught],
        },
        "settings": get_settings_map(db),
        "books": books,                                    # V2-A02 ✓ کتاب‌ها
        "test_sessions": [
            {
                "id": s.id, "mode": s.mode, "label": s.label, "source": s.source,
                "resource_id": s.resource_id, "resource_title": s.resource_title, "exam_id": s.exam_id,
                "total_count": s.total_count, "correct_count": s.correct_count, "wrong_count": s.wrong_count,
                "unanswered_count": s.unanswered_count, "not_entered_count": s.not_entered_count,
                "percent_konkur": s.percent_konkur, "percent_no_penalty": s.percent_no_penalty,
                "penalty_k": s.penalty_k, "planned_duration": s.planned_duration,
                "actual_duration": s.actual_duration,
                "started_at": _iso(s.started_at), "finished_at": _iso(s.finished_at),
            }
            for s in sessions
        ],
        "attempts": [                                       # V2-A02 ✓ attemptها
            {
                "id": a.id, "session_id": a.session_id, "question_id": a.question_id,
                "topic_id": a.topic_id, "question_number": a.question_number, "topic_title": a.topic_title,
                "status": a.status, "result": a.result, "answer": a.answer, "correct_answer": a.correct_answer,
                "answer_key_version": a.answer_key_version, "duration_seconds": a.duration_seconds,
                "solved_at": _iso(a.solved_at),
            }
            for a in attempts
        ],
        "error_notes": [
            {
                "id": n.id, "book_title": n.book_title, "topic_title": n.topic_title,
                "question_number": n.question_number, "your_answer": n.your_answer,
                "correct_answer": n.correct_answer, "error_type": n.error_type, "note": n.note,
            }
            for n in notes
        ],
        "question_marks": [
            {"question_id": m.question_id, "review": m.review, "important": m.important, "hard": m.hard}
            for m in marks
        ],
        "review_items": [
            {
                "id": r.id, "question_id": r.question_id, "topic_title": r.topic_title,
                "book_title": r.book_title, "source": r.source, "status": r.status, "critical": r.critical,
                "wrong_count": r.wrong_count, "review_count": r.review_count, "cycle_index": r.cycle_index,
                "scheduled_date": _iso(r.scheduled_date), "last_reviewed_at": _iso(r.last_reviewed_at),
            }
            for r in review_items
        ],
        "learning_states": [
            {
                "topic_id": s.topic_id, "coverage": s.coverage, "accuracy": s.accuracy,
                "retention_est": s.retention_est, "exam_readiness": s.exam_readiness,
                "confidence": s.confidence, "weakness": s.weakness,
            }
            for s in states
        ],
        "goals": [
            {"id": g.id, "title": g.title, "kind": g.kind, "target_date": _iso(g.target_date)} for g in goals
        ],
        "exams": [
            {
                "id": e.id, "title": e.title, "kind": e.kind, "status": e.status,
                "scheduled_date": _iso(e.scheduled_date), "subjects": e.subjects,
                "planned_duration_minutes": e.planned_duration_minutes,
                "actual_duration_seconds": e.actual_duration_seconds, "scoring": e.scoring,
            }
            for e in exams
        ],
        "plan_tasks": [
            {
                "id": t.id, "date": _iso(t.date), "kind": t.kind, "title": t.title,
                "topic_title": t.topic_title, "minutes": t.minutes, "count": t.count,
                "status": t.status, "locked": t.locked, "source": t.source, "reason_code": t.reason_code,
            }
            for t in plan_tasks
        ],
        "time_blocks": [
            {
                "date": _iso(b.date), "kind": b.kind, "start_minutes": b.start_minutes,
                "end_minutes": b.end_minutes, "title": b.title, "source": b.source,
            }
            for b in time_blocks
        ],
        "capacity_snapshots": [
            {
                "date": _iso(s.date), "school_minutes": s.school_minutes,
                "available_study_minutes": s.available_study_minutes,
                "estimated_capacity_tasks": s.estimated_capacity_tasks,
                "completion_rate": s.completion_rate, "source": s.source,
            }
            for s in capacity
        ],
        "recommendations": [
            {
                "id": r.id, "date": _iso(r.date), "status": r.status,
                "payload": r.payload, "reasons": r.reasons,
            }
            for r in recs
        ],
        "checkins": [
            {
                "date": _iso(ci.date), "energy": ci.energy, "focus": ci.focus, "motivation": ci.motivation,
                "stress": ci.stress, "fatigue": ci.fatigue,
            }
            for ci in checkins
        ],
    }


# --- Excel (openpyxl، RTL) ---------------------------------------------------------------------

def excel_bytes(db: Session, student: Student) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = full_json(db, student)
    ov = analytics_overview(db, student, days=30)

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", name="Vazirmatn")
    header_fill = PatternFill("solid", fgColor="4F6BED")
    rtl = Alignment(horizontal="right", vertical="center")

    def sheet(name: str, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None):
        ws = wb.create_sheet(name)
        ws.sheet_view.rightToLeft = True  # RTL کامل برای فارسی
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = rtl
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, w in enumerate(widths or [18] * len(headers), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    # خلاصه — سه بلوک متریک همیشه جدا (V2-A01 حتی در خروجی)
    ws = wb.active
    ws.title = "خلاصه"
    ws.sheet_view.rightToLeft = True
    ws.append([f"گزارش ALEMS — {data['profile']['email']} — {_jalali_str(_today())}"])
    ws["A1"].font = Font(bold=True, size=13, name="Vazirmatn")
    cov, acc, vol = ov["coverage"], ov["accuracy"], ov["volume"]
    ws.append([])
    ws.append(["۱) پوشش (Coverage)"])
    ws.append(["موضوع دیده‌شده", f"{cov['topics_attempted']} از {cov['topics_total']}"])
    ws.append(["سوال دیده‌شده", f"{cov['questions_attempted']} از {cov['questions_total']}"])
    ws.append([])
    ws.append(["۲) دقت (Accuracy) — ۳۰ روز اخیر"])
    ws.append(["درست / غلط", f"{acc['correct']} / {acc['wrong']}"])
    ws.append(["دقت پاسخ‌داده‌ها", acc["answered_accuracy"]])
    ws.append(["درصد کنکوری (با جریمه)", acc["percent_konkur"]])
    ws.append(["درصد بدون جریمه", acc["percent_no_penalty"]])
    ws.append([])
    ws.append(["۳) حجم (Volume) — ۳۰ روز اخیر"])
    ws.append(["تعداد پاسخ", vol["attempts"]])
    ws.append(["جلسه تست", vol["sessions"]])
    ws.append(["دقیقه مطالعه", vol["study_minutes"]])
    ws.append(["روز فعال", vol["active_days"]])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22

    sheet(
        "کتاب‌ها",
        ["کتاب", "ناشر", "درس", "فصل/موضوع", "نوع", "تعداد سوال"],
        [
            [b["title"], b["publisher"], b["subject"] or "", t["title"], t["block_type"], len(t["questions"])]
            for b in data["books"]
            for t in _flat_topics(b["tree"])
        ],
        [26, 16, 12, 30, 12, 10],
    )
    sheet(
        "تست‌ها",
        ["برچسب", "کتاب", "تاریخ شروع", "کل", "درست", "غلط", "نزده", "درصد کنکوری", "درصد بدون جریمه"],
        [
            [s["label"] or s["mode"], s["resource_title"] or "", _iso(s["started_at"]) or "",
             s["total_count"], s["correct_count"], s["wrong_count"], s["unanswered_count"],
             s["percent_konkur"], s["percent_no_penalty"]]
            for s in data["test_sessions"]
        ],
        [22, 22, 20, 8, 8, 8, 8, 14, 16],
    )
    sheet(
        "تلاش‌ها",
        ["زمان", "موضوع", "شماره سوال", "وضعیت", "نتیجه", "پاسخ شما", "پاسخ درست", "ثانیه"],
        [
            [a["solved_at"] or "", a["topic_title"] or "", a["question_number"], a["status"],
             a["result"], a["answer"] or "", a["correct_answer"] or "", a["duration_seconds"] or ""]
            for a in data["attempts"]
        ],
        [22, 26, 10, 12, 10, 10, 10, 8],
    )
    sheet(
        "مرور",
        ["موضوع", "کتاب", "منبع", "وضعیت", "بحرانی", "غلط‌ها", "مرورها", "سررسید"],
        [
            [r["topic_title"] or "", r["book_title"] or "", r["source"], r["status"],
             "بله" if r["critical"] else "", r["wrong_count"], r["review_count"], r["scheduled_date"] or ""]
            for r in data["review_items"]
        ],
        [26, 22, 12, 10, 8, 8, 8, 12],
    )
    sheet(
        "برنامه",
        ["تاریخ", "عنوان", "نوع", "دقیقه", "وضعیت", "قفل", "منبع"],
        [
            [t["date"], t["title"], t["kind"], t["minutes"], t["status"],
             "بله" if t["locked"] else "", t["source"]]
            for t in data["plan_tasks"]
        ],
        [12, 34, 10, 8, 10, 6, 10],
    )
    sheet(
        "آزمون‌ها",
        ["عنوان", "نوع", "وضعیت", "تاریخ", "درصد کنکوری", "درصد بدون جریمه"],
        [
            [e["title"], e["kind"], e["status"], e["scheduled_date"] or "",
             (e["scoring"] or {}).get("percent_konkur"), (e["scoring"] or {}).get("percent_no_penalty")]
            for e in data["exams"]
        ],
        [26, 12, 12, 12, 14, 16],
    )

    buf = io.BytesIO()
    wb.save(buf)
    logger.info("excel export built bytes=%d sheets=%d", buf.tell(), len(wb.sheetnames))
    return buf.getvalue()


def _flat_topics(tree: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for node in tree:
        out.append(node)
        out.extend(_flat_topics(node.get("children") or []))
    return out


# --- PDF فارسی RTL (reportlab + Vazirmatn + reshaper/bidi) ---------------------------------------

_pdf_fonts_ready = False


def _ensure_fonts() -> None:
    global _pdf_fonts_ready
    if _pdf_fonts_ready:
        return
    import os

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    base = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "static", "fonts")
    pdfmetrics.registerFont(TTFont("Vazirmatn", os.path.join(base, "Vazirmatn-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("Vazirmatn-Bold", os.path.join(base, "Vazirmatn-Bold.ttf")))
    _pdf_fonts_ready = True


def _fa(s: Any) -> str:
    """shape + bidi — متن فارسی درست در PDF."""
    import arabic_reshaper
    from bidi.algorithm import get_display

    return get_display(arabic_reshaper.reshape(str(s)))


class _PdfDoc:
    """canvas ساده RTL با صفحه‌بندی خودکار."""

    def __init__(self) -> None:
        from reportlab.pdfgen import canvas as rl_canvas

        self.buf = io.BytesIO()
        self.c = rl_canvas.Canvas(self.buf, pagesize=(595.27, 841.89))  # A4
        self.y = 790.0
        self.right = 545.0
        self.left = 50.0

    def _check(self, need: float = 18) -> None:
        if self.y - need < 55:
            self.c.showPage()
            self.y = 790.0

    def title(self, text: str, size: int = 15) -> None:
        self._check(30)
        self.c.setFont("Vazirmatn-Bold", size)
        self.c.drawRightString(self.right, self.y, _fa(text))
        self.y -= size + 10

    def line(self, text: str, size: int = 10, indent: float = 0) -> None:
        self._check()
        self.c.setFont("Vazirmatn", size)
        self.c.drawRightString(self.right - indent, self.y, _fa(text))
        self.y -= size + 7

    def kv(self, key: str, value: Any, indent: float = 0) -> None:
        self.line(f"{key}: {value}", indent=indent)

    def section(self, text: str) -> None:
        self._check(26)
        self.y -= 4
        self.c.setFont("Vazirmatn-Bold", 12)
        self.c.drawRightString(self.right, self.y, _fa(text))
        self.y -= 16
        self.c.setStrokeColorRGB(0.8, 0.8, 0.85)
        self.c.line(self.left, self.y + 6, self.right, self.y + 6)

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
        """جدول RTL — اولین ستون سمت راست."""
        self._check(24)
        x = self.right
        self.c.setFont("Vazirmatn-Bold", 9)
        for h, w in zip(headers, widths):
            self.c.drawRightString(x, self.y, _fa(h))
            x -= w
        self.y -= 6
        self.c.setStrokeColorRGB(0.75, 0.75, 0.8)
        self.c.line(self.left, self.y, self.right, self.y)
        self.y -= 12
        self.c.setFont("Vazirmatn", 9)
        for row in rows:
            self._check(14)
            x = self.right
            for cell, w in zip(row, widths):
                self.c.drawRightString(x, self.y, _fa(cell))
                x -= w
            self.y -= 13

    def save(self) -> bytes:
        self.c.setFont("Vazirmatn", 8)
        self.c.drawCentredString(297, 30, _fa("ساخته‌شده با ALEMS — گزارش شخصی دانش‌آموز"))
        self.c.save()
        return self.buf.getvalue()


def _metrics_sections(doc: _PdfDoc, rep: dict[str, Any]) -> None:
    cov, acc, vol = rep["coverage"], rep["accuracy"], rep["volume"]
    doc.section("۱) پوشش")
    doc.kv("موضوع دیده‌شده", f"{_num(cov['topics_attempted'])} از {_num(cov['topics_total'])} ({_pct(cov['topics_ratio'])})")
    doc.kv("سوال دیده‌شده", f"{_num(cov['questions_attempted'])} از {_num(cov['questions_total'])} ({_pct(cov['questions_ratio'])})")
    doc.section("۲) دقت")
    doc.kv("درست / غلط", f"{_num(acc['correct'])} / {_num(acc['wrong'])}")
    doc.kv("دقت روی پاسخ‌داده‌ها", _pct(acc["answered_accuracy"]))
    doc.kv("درصد کنکوری (با جریمه)", "—" if acc["percent_konkur"] is None else _num(f"{acc['percent_konkur']}٪"))
    doc.kv("درصد بدون جریمه", "—" if acc["percent_no_penalty"] is None else _num(f"{acc['percent_no_penalty']}٪"))
    doc.section("۳) حجم")
    doc.kv("تعداد پاسخ", _num(vol["attempts"]))
    doc.kv("جلسه تست", _num(vol["sessions"]))
    doc.kv("دقیقه مطالعه", _num(vol["study_minutes"]))
    doc.kv("روز فعال", _num(vol["active_days"]))
    doc.kv("مرور انجام‌شده", _num(vol["reviews_done"]))


def pdf_bytes(db: Session, student: Student, report_kind: str = "weekly", when: str | None = None) -> bytes:
    _ensure_fonts()
    user = db.execute(select(User).where(User.id == student.user_id)).scalar_one_or_none()
    doc = _PdfDoc()

    if report_kind == "daily":
        rep = report_service.daily(db, student, when)
        doc.title(f"گزارش روزانه — {rep['weekday_fa']} {_jalali_str(dt.date.fromisoformat(rep['date']))}")
    elif report_kind == "monthly":
        rep = report_service.monthly(db, student, when)
        doc.title(f"گزارش ماهانه — {rep['month']}")
    elif report_kind == "summary":
        rep = None
        doc.title("خلاصه تحلیل — ۳۰ روز اخیر")
    else:
        report_kind = "weekly"
        rep = report_service.weekly(db, student, when)
        doc.title(f"گزارش هفتگی — هفته {_jalali_str(dt.date.fromisoformat(rep['week_start']))}")

    doc.line(f"دانش‌آموز: {user.email if user else '—'} · تاریخ ساخت: {_jalali_str(_today())}", size=9)

    if report_kind == "summary":
        ov = analytics_overview(db, student, days=30)
        _metrics_sections(doc, ov)
        if ov["weaknesses"]:
            doc.section("نقاط ضعف")
            doc.table(
                ["موضوع", "آمادگی آزمون", "دقت", "پوشش"],
                [
                    [w["topic_title"] or "—", _num(round(w["exam_readiness"], 2)),
                     _pct(w["accuracy"]), _pct(w["coverage"])]
                    for w in ov["weaknesses"]
                ],
                [220, 100, 90, 90],
            )
        tgt = ov["konkurs_target"]
        if tgt["has_target"]:
            doc.section("هدف کنکور (نمایش کیفی)")
            doc.line(tgt["message_fa"], size=10)
    else:
        _metrics_sections(doc, rep)
        if report_kind == "weekly":
            doc.section("روزهای هفته")
            doc.table(
                ["روز", "تاریخ", "تست", "درست", "غلط", "دقیقه مطالعه", "کار انجام/کل"],
                [
                    [d["weekday_fa"], d["date_jalali"], _num(d["attempts"]), _num(d["correct"]),
                     _num(d["wrong"]), _num(d["study_minutes"]),
                     f"{_num(d['tasks_done'])}/{_num(d['tasks_total'])}"]
                    for d in rep["days"]
                ],
                [70, 90, 55, 55, 55, 95, 80],
            )
            prev = rep["previous_week"]
            doc.section("مقایسه با هفته قبل")
            doc.kv("تغییر تعداد تست", f"{_num(prev['attempts_delta'])}+")
            doc.kv("تغییر دقیقه مطالعه", f"{_num(prev['study_minutes_delta'])}+")
        elif report_kind == "monthly":
            doc.section("هفته‌های ماه")
            doc.table(
                ["شروع هفته", "روزها", "تست", "درست", "غلط", "دقیقه مطالعه"],
                [
                    [w["week_start_jalali"], _num(w["days"]), _num(w["attempts"]), _num(w["correct"]),
                     _num(w["wrong"]), _num(w["study_minutes"])]
                    for w in rep["weeks"]
                ],
                [110, 60, 70, 70, 70, 110],
            )
        if rep.get("top_topics"):
            doc.section("پرتکرارترین موضوع‌ها")
            doc.table(
                ["موضوع", "تعداد پاسخ"],
                [[t["topic_title"] or "—", _num(t["attempts"])] for t in rep["top_topics"]],
                [300, 100],
            )
        if rep.get("exams"):
            doc.section("آزمون‌ها")
            doc.table(
                ["عنوان", "تاریخ", "وضعیت", "درصد کنکوری", "بدون جریمه"],
                [
                    [e["title"], e["scheduled_date_jalali"], e["status"],
                     "—" if e["percent_konkur"] is None else _num(f"{e['percent_konkur']}٪"),
                     "—" if e["percent_no_penalty"] is None else _num(f"{e['percent_no_penalty']}٪")]
                    for e in rep["exams"]
                ],
                [160, 90, 80, 90, 90],
            )

    out = doc.save()
    logger.info("pdf export built kind=%s bytes=%d", report_kind, len(out))
    return out
